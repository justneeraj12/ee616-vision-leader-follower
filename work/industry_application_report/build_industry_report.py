from __future__ import annotations

import json
import shutil
import sys
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'work' / 'industry_application_report'
OUTPUT = ROOT / 'deliverables' / 'Neeraj_Kanchani_EE616_Industry_Application_and_System_Rationale.docx'
REFERENCE = ROOT / 'deliverables' / 'Neeraj_Kanchani_EE616_Project_Update_Package.docx'
SIM_SUMMARY = ROOT / 'work' / 'toy_simulation' / 'warehouse_results' / 'warehouse_multi_summary.json'
SIM_FRAME = ROOT / 'work' / 'toy_simulation' / 'warehouse-video-qa' / 'frame-03.png'
DIAGRAM = WORK / 'industry_architecture.png'

sys.path.insert(0, str(ROOT / 'work' / 'professor_package'))
from build_package import (  # noqa: E402
    BLACK,
    BLUE,
    LIGHT_BLUE,
    MID,
    NAVY,
    PALE,
    WHITE,
    add_heading,
    add_hyperlink,
    add_para,
    add_table,
    clear_document_body,
    configure_styles,
    set_cell_margins,
    set_repeat_table_header,
    set_run_font,
)


def keep_table_rows(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)


def add_figure(doc, path: Path, caption: str, alt: str, width: float = 7.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('name', path.stem)
    shape._inline.docPr.set('descr', alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_run_font(cap.add_run(caption), size=8.7, italic=True, color=MID)


def add_reference(doc, number: int, citation: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(f'[{number}] {citation} '), size=8.8)
    add_hyperlink(p, 'Source link', url)


def make_architecture_diagram() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    width, height = 3000, 1560
    img = Image.new('RGB', (width, height), '#F5F9FC')
    draw = ImageDraw.Draw(img)
    regular = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 27)
    small = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 23)
    bold = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 31)
    title = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 49)

    def box(x, y, w, h, heading, body, fill='#FFFFFF', edge='#0B3153'):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=22, fill=fill, outline=edge, width=5)
        draw.text((x + 28, y + 24), heading, font=bold, fill='#0B3153')
        lines = []
        for logical in body.split('\n'):
            lines.extend(textwrap.wrap(logical, width=max(16, int(w / 24))))
        draw.multiline_text((x + 28, y + 78), '\n'.join(lines), font=small, fill='#34495E', spacing=8)

    def arrow(x1, y1, x2, y2, color='#2F6F9F'):
        draw.line((x1, y1, x2, y2), fill=color, width=7)
        import math
        a = math.atan2(y2 - y1, x2 - x1)
        size = 28
        spread = 0.55
        draw.polygon([
            (x2, y2),
            (x2 - size * math.cos(a - spread), y2 - size * math.sin(a - spread)),
            (x2 - size * math.cos(a + spread), y2 - size * math.sin(a + spread)),
        ], fill=color)

    draw.rounded_rectangle((32, 28, width - 32, 150), radius=16, fill='#0B3153')
    draw.text((92, 62), 'Proposed Warehouse Material Delivery Convoy', font=title, fill='white')

    box(70, 250, 520, 270, 'WMS or fleet manager', 'Assigns destination, route, and delivery job to the leader. Receives status and logs.', '#E5F0F8')
    box(820, 225, 650, 320, 'Leader robot', 'Owns route intent and global localization. Plans aisle motion, avoids obstacles, enforces the speed profile, and exposes a stable visual target.')
    box(1720, 225, 500, 320, 'Follower 1', 'Carries material. Uses its own forward camera and odometry to track the leader with a commanded gap.')
    box(2440, 225, 500, 320, 'Follower 2 and beyond', 'Repeats the same local behavior by tracking the robot immediately ahead. No central low-level motion command.')
    arrow(590, 385, 820, 385)
    arrow(1470, 385, 1720, 385)
    arrow(2220, 385, 2440, 385)

    draw.text((92, 660), 'Independent functions on every robot', font=bold, fill='#0B3153')
    box(70, 730, 650, 250, 'Formation coordination', 'Camera detection, relative range and bearing, state estimation, bounded control, and recovery states.', '#FFFFFF')
    box(835, 730, 650, 250, 'Safety layer', 'Safety-rated person and obstacle detection, emergency stop, safe braking, speed and zone supervision.', '#FFF6E8', '#A76D10')
    box(1600, 730, 650, 250, 'Operations layer', 'Battery health, diagnostics, cybersecurity, fault records, maintenance, and dispatch integration.', '#FFFFFF')
    box(2365, 730, 565, 250, 'Evidence layer', 'Versioned configuration, time-synchronized logs, fault injection, and repeatable acceptance tests.', '#FFFFFF')

    draw.rounded_rectangle((70, 1115, 2930, 1435), radius=20, fill='#EAF2F7', outline='#60758A', width=4)
    draw.text((105, 1150), 'System boundary for this EE 616 project', font=bold, fill='#0B3153')
    boundary = (
        'I will study the follower formation-coordination layer in simulation. The project does not claim that a monocular camera '
        'is a complete industrial safety system. A deployment would require an independent safety architecture, verified stopping '
        'behavior, site risk assessment, commissioning, and compliance with ISO 3691-4 and applicable local requirements.'
    )
    wrapped = textwrap.wrap(boundary, width=145)
    draw.multiline_text((105, 1210), '\n'.join(wrapped), font=regular, fill='#34495E', spacing=10)
    img.save(DIAGRAM, quality=95)


def build() -> None:
    with SIM_SUMMARY.open(encoding='utf-8') as stream:
        result = json.load(stream)
    make_architecture_diagram()
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_document_body(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)

    top = doc.add_paragraph()
    top.paragraph_format.space_before = Pt(82)
    top.paragraph_format.space_after = Pt(4)
    set_run_font(top.add_run('EE 616 Application Rationale'), size=21, color=MID)

    title = doc.add_paragraph(style='Title')
    title.add_run('Industry Application and System Rationale for Vision Only Leader Follower Warehouse Coordination')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(12)
    for run in title.runs:
        set_run_font(run, size=28, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(58)
    set_run_font(subtitle.add_run('Proposed flexible material delivery convoy for faculty review'), size=13.5, color='34495E')

    meta = doc.add_table(rows=1, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_repeat_table_header(meta.rows[0])
    for index, (label, value) in enumerate([
        ('STATUS', 'Draft for faculty review'),
        ('AUTHOR', 'Neeraj Kumar Kanchani'),
        ('UPDATED', 'September 17, 2026'),
    ]):
        cell = meta.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.width = Inches(2.35)
        set_cell_margins(cell, 85, 45, 85, 45)
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        set_run_font(p1.add_run(label), size=8.5, bold=True, color=MID)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=10.1, color='24384D')

    doc.add_paragraph().paragraph_format.space_after = Pt(5)
    t = add_table(doc, ['Field', 'Details'], [
        ('Faculty reviewer', 'Professor Masudul Imtiaz'),
        ('Course', 'EE 616-57, section 9189, four credits'),
        ('Proposed application', 'Flexible autonomous material-delivery convoy for warehouse and manufacturing aisles'),
        ('Research boundary', 'Simulation study of decentralized visual following on constrained edge hardware'),
        ('Review requested', 'Confirm the application framing and approve the ROS 2 and Gazebo implementation scope'),
    ], widths=[1.55, 5.45], font_size=9.0)
    keep_table_rows(t)

    doc.add_page_break()
    add_heading(doc, '1 Proposal and Review Request')
    add_para(doc,
        'I propose to develop and evaluate a leader-follower robot convoy for repeated material movement in warehouses and manufacturing facilities. The leader receives the delivery mission and owns the route. Each follower carries material and uses a forward RGB camera, together with its own wheel odometry, to maintain a safe commanded distance from the robot immediately ahead. The research contribution is the decentralized formation-control layer: relative visual perception, state estimation, bounded control, loss-of-view recovery, and measured behavior on limited edge hardware.')
    add_para(doc,
        'The industrial value is a flexible alternative to moving one cart per trip or mechanically coupling every load into a fixed tugger train. A site could add or remove follower carts as payload demand changes while keeping route planning concentrated on the leader. This idea is consistent with existing warehouse practices, including tugger trains for multi-cart movement and commercial follow-me autonomous mobile robots for picking and internal transport. My project will test a narrower question: whether a camera-based follower can maintain an ordered convoy through warehouse-like paths under controlled simulation conditions.')
    add_para(doc,
        'I am asking for feedback on the application definition, the separation between coordination and safety, and the planned validation path. I am not presenting the current toy results as Gazebo evidence or as proof of deployment safety.')

    add_heading(doc, '2 Industrial Problem')
    add_para(doc,
        'Warehouses and manufacturing cells repeatedly move parts, picked goods, empty containers, and work-in-process inventory between receiving, storage, packing, and line-side stations. These movements consume operator time and create traffic in narrow aisles. Toyota describes tugger trains as a way to move multiple carts in fewer trips and support lean or just-in-time delivery. Zebra lists order picking, post-pick transport, receiving and putaway, and replenishment among current autonomous mobile robot workflows. These are established material-handling problems rather than a hypothetical use case. [2][3]')
    add_para(doc,
        'The proposed convoy addresses jobs where several loads should travel along the same route, but a facility would benefit from changing the number of carts without rebuilding a mechanical train or programming a complete route on every unit. The leader carries the route and navigation capability. Followers reproduce the local motion of the robot ahead while maintaining separation. The convoy can therefore be configured for a small delivery during normal operation and expanded for peak demand, subject to aisle geometry and site safety limits.')

    h = add_heading(doc, '3 Industry Evidence and Comparable Systems')
    h.paragraph_format.page_break_before = True
    add_para(doc,
        'The closest commercial precedent is the follow-me autonomous mobile robot. DHL reported testing the EffiBOT trolley in warehouse operations. The robot followed employees, carried loads, avoided obstacles, and could travel to a drop-off point. DHL later described follow-me and point-to-point operation as ways to reduce repetitive walking and material movement. Zebra also describes a follow-me workflow built around one picker and one robot. These systems establish that local following behavior has practical value, while also showing that current products often use a single-person or single-robot relationship rather than a multi-robot visual convoy. [1][4][5]')
    t = add_table(doc, ['Industrial pattern', 'What industry already uses', 'What this project adds or studies'], [
        ('Follow-me AMR', 'One robot follows a worker during picking or transport.', 'Robot-to-robot visual following without human motion commands.'),
        ('Tugger train', 'One powered vehicle pulls several mechanically coupled carts.', 'Software-defined spacing and the ability to add or remove independently driven carts.'),
        ('Fleet-managed AMRs', 'A fleet manager dispatches independent robots and manages traffic.', 'A convoy whose low-level formation remains local even if mission assignment is centralized.'),
        ('Fixed automation', 'Conveyors or guided routes provide predictable material flow.', 'A reconfigurable route-following concept for existing aisles where fixed infrastructure may be undesirable.'),
    ], widths=[1.25, 2.65, 3.10], font_size=8.7)
    keep_table_rows(t)
    add_para(doc,
        'I therefore view the project as an applied research bridge between follow-me assistance, tugger-train logistics, and multi-robot autonomy. It does not assume that warehouses currently deploy this exact camera-only chain at scale. The report uses published commercial systems to establish the operational need, then treats the proposed convoy as the research contribution.')

    doc.add_page_break()
    add_heading(doc, '4 Proposed Operating Scenario')
    add_figure(doc, DIAGRAM,
        'Figure 1  Proposed operating model and the boundary of the EE 616 research system',
        'A warehouse management system assigns a mission to a leader robot. The leader is followed by a chain of load-carrying robots. Every robot has independent coordination, safety, operations, and evidence functions. A note distinguishes the simulation research layer from deployment safety requirements.', 7.0)
    add_para(doc,
        'A warehouse management system or fleet manager assigns a replenishment or transport job to the leader. The leader travels from a pickup point through warehouse aisles to a destination. Follower 1 observes the leader. Follower 2 observes Follower 1, and the pattern continues for additional carts. This predecessor-following topology keeps each perception problem local and avoids requiring the last cart to see the leader through the entire convoy.')

    leader_heading = add_heading(doc, '5 Leader Responsibilities')
    leader_heading.paragraph_format.page_break_before = True
    add_para(doc,
        'The leader is the route owner and the highest-capability robot in the convoy. In the research simulation, it follows a defined waypoint route. In an industrial version, it would receive a mission from the warehouse or fleet system and use global localization, route planning, obstacle avoidance, and traffic rules to execute that mission. It would also expose a stable rear visual target or distinctive geometry that followers can detect reliably.')
    t = add_table(doc, ['Leader function', 'Reason it belongs on the leader'], [
        ('Mission and destination', 'One route owner prevents every follower from independently interpreting the job.'),
        ('Global localization and path planning', 'Only the leader needs the complete route to guide the convoy through the facility.'),
        ('Speed and curvature profile', 'A conservative leader profile limits acceleration and cornering demands passed down the chain.'),
        ('Obstacle response', 'The leader initiates planned slowing or stopping when the route becomes blocked.'),
        ('Trackable rear target', 'A known appearance and geometry improve range, bearing, and confidence estimation for Follower 1.'),
        ('Convoy status', 'The leader should not continue a mission when a follower reports a safety stop or unrecoverable fault.'),
    ], widths=[2.0, 5.0], font_size=8.8)
    keep_table_rows(t)
    add_para(doc,
        'The leader remains a mission-level dependency, so loss of the leader must stop or reassign the convoy. Decentralization does not remove that operational dependency. It removes the need for a remote computer to generate every follower wheel command.')

    doc.add_page_break()
    add_heading(doc, '6 Follower Responsibilities and Local Control')
    add_para(doc,
        'Each follower is an independently driven load carrier. It detects the robot immediately ahead, estimates relative range and bearing, predicts short measurement gaps, and commands bounded linear and angular velocity. The initial design uses TRACK, PREDICT, and SAFE STOP states. A follower enters prediction only for a short, bounded interval and stops when confidence or uncertainty crosses a declared limit.')
    t = add_table(doc, ['Follower input', 'Local computation', 'Required behavior'], [
        ('Forward RGB image', 'Leader or predecessor detection and confidence', 'Use only recent, valid observations inside the declared field of view.'),
        ('Wheel odometry', 'Motion prediction between image updates', 'Reject stale timestamps and expose uncertainty growth.'),
        ('Relative range and bearing', 'State estimation and formation error', 'Maintain the commanded gap without exceeding acceleration or speed limits.'),
        ('Detection health', 'Supervisory state transition', 'Track when confident, predict briefly, and stop safely after the timeout.'),
        ('Predecessor and convoy status', 'Fault propagation', 'Prevent trailing robots from continuing into a stopped section of the convoy.'),
    ], widths=[1.55, 2.55, 2.90], font_size=8.7)
    keep_table_rows(t)

    add_heading(doc, '7 Why Leader Follower Coordination Is Useful')
    add_para(doc,
        'Leader-follower coordination is useful when several loads share a route and the facility wants to change convoy capacity without mechanically coupling every cart. The leader concentrates global navigation and traffic decisions. Followers reuse the motion intent expressed by the robot ahead. This can reduce duplicated route-planning hardware and simplify job configuration, although the final design must retain enough sensing and compute on every robot for independent safety and fault handling.')
    t = add_table(doc, ['Benefit', 'Operational reason', 'Condition for the benefit to hold'], [
        ('Variable carrying capacity', 'Followers can be added or removed for a delivery wave.', 'The route and stopping model must be validated for the resulting convoy length.'),
        ('Fewer repeated trips', 'One mission can move several powered carts along the same path.', 'Pickup and drop-off handling must support the convoy workflow.'),
        ('Brownfield fit', 'The concept uses mobile robots in existing aisles instead of a new fixed conveyor.', 'The site must have sufficient aisle width, visibility, floor quality, and safe zones.'),
        ('Local low-level control', 'Followers do not depend on continuous remote wheel commands.', 'Status communication, fault propagation, and mission supervision are still required.'),
        ('Staged development', 'The formation layer can be tested first in simulation and then on one pair.', 'Each stage must have measurable acceptance criteria before scaling the chain.'),
    ], widths=[1.35, 2.65, 3.00], font_size=8.5)
    keep_table_rows(t)

    doc.add_page_break()
    add_heading(doc, '8 Why Vision Only Following Is Worth Studying')
    add_para(doc,
        'A forward RGB camera is low cost, compact, and information-rich. It can estimate target identity, image position, and apparent scale with one sensor. For this project, camera-only coordination creates a focused research problem that fits the available RTX 3050 Ti laptop GPU and can be measured in Gazebo. It also exposes a real control issue: the follower must keep the target visible while maintaining spacing. Visibility therefore becomes part of the motion constraint rather than a separate perception detail.')
    add_para(doc,
        'The same choice creates important limits. Monocular scale is ambiguous, appearance changes with lighting and viewpoint, sharp corners can cause occlusion, and one robot can hide another. A long predecessor-following chain can accumulate spacing and steering error. The project will study these failure modes instead of assuming that a camera replaces safety-rated ranging or person detection.')

    add_heading(doc, '9 Current Toy Simulation Evidence')
    add_figure(doc, SIM_FRAME,
        'Figure 2  Representative frame from the current three-follower warehouse toy simulation',
        'A top-down warehouse maze contains shelf blocks, a leader, and three followers moving along a multi-segment route. Trails show the convoy path and labels identify each robot.', 6.8)
    rmse = result['spacing_rmse_m']
    add_para(doc,
        f'The current two-dimensional kinematic run contains one leader, {result["followers"]} followers, six shelf blocks, and '
        f'{result["completed_waypoints"]} path segments over {result["duration_s"]:.0f} simulated seconds. It recorded zero collision samples. '
        f'Spacing RMSE was {rmse["follower_1"]:.3f} m for Follower 1, {rmse["follower_2"]:.3f} m for Follower 2, and '
        f'{rmse["follower_3"]:.3f} m for Follower 3. None of the followers entered SAFE STOP during this nominal run.')
    add_para(doc,
        'This result demonstrates the planned chain topology, warehouse route, logging, and video presentation. It does not include ROS 2 message timing, Gazebo physics, rendered camera images, learned detection, wheel slip, person traffic, or a safety-rated obstacle system. The increase in spacing error toward the rear of the convoy is also a warning that multi-follower stability must be measured rather than assumed.')

    doc.add_page_break()
    add_heading(doc, '10 Industry Level Architecture')
    add_para(doc,
        'A credible industrial system needs more than the formation controller. I would separate the system into five layers so that a failure in visual following does not remove the robot safety functions.')
    t = add_table(doc, ['Layer', 'Primary responsibility', 'Evidence required before deployment'], [
        ('Mission and fleet', 'Job assignment, route ownership, traffic coordination, charging, and recovery.', 'Interface tests, route permissions, fault recovery, and dispatch records.'),
        ('Leader autonomy', 'Localization, planning, obstacle response, and convoy-aware speed profile.', 'Localization coverage, stopping-distance tests, route and intersection cases.'),
        ('Follower formation', 'Visual detection, relative estimation, bounded following, and loss-of-view handling.', 'Error distributions, latency, field-of-view limits, occlusion tests, and chain stability.'),
        ('Independent safety', 'Person and obstacle detection, emergency stop, safe braking, speed and zone supervision.', 'Safety functions designed and validated under the applicable industrial standard.'),
        ('Operations and assurance', 'Diagnostics, cybersecurity, maintenance, configuration control, and incident logs.', 'Versioned releases, health monitoring, access control, and commissioning records.'),
    ], widths=[1.35, 3.15, 2.50], font_size=8.4)
    keep_table_rows(t)

    add_heading(doc, '11 Safety and Production Readiness Boundary')
    add_para(doc,
        'ISO 3691-4 covers safety requirements for driverless industrial trucks, including automated guided vehicles, autonomous mobile robots, carts, and tuggers. OSHA also identifies struck-by and caught-between hazards when robots and automated equipment operate in warehouses. These sources mean that collision-free simulation is necessary evidence for development, but it is not a safety case. [6][7]')
    add_para(doc,
        'A deployable convoy would require an independent safety-rated layer on every robot, emergency-stop access, verified braking and stopping distance, overspeed protection, person and obstacle detection, communications supervision, safe behavior after faults, audible or visual warnings, site risk assessment, cybersecurity, commissioning, and documented maintenance. The formation camera could contribute operational awareness, but I would not use it as the sole safety sensor.')
    add_para(doc,
        'Gazebo will be used to find control defects, measure nominal and adverse behavior, and reproduce failure cases before physical testing. It will not be described as certification. Any physical prototype would begin at low speed in a controlled area with a separate safety plan and faculty approval.')

    doc.add_page_break()
    add_heading(doc, '12 Evaluation Plan')
    t = add_table(doc, ['Question', 'Experiment', 'Decision measure'], [
        ('Can one follower estimate the predecessor state?', 'Range, bearing, heading, lighting, and off-axis camera sweeps in Gazebo.', 'Range and bearing error distributions, confidence calibration, and valid operating envelope.'),
        ('Can one pair maintain formation?', 'Straight, curved, stop-and-go, and speed-change trajectories.', 'Spacing error, yaw error, overshoot, control timing, and minimum separation.'),
        ('Can the system recover from loss of view?', 'Short occlusions, dropped frames, blur, and cornering cases.', 'Time to reacquire, SAFE STOP rate, and distance traveled without a valid observation.'),
        ('Does the chain remain stable?', 'One, two, and three followers on the same routes and disturbance cases.', 'Error growth by follower index, collision count, and string-stability indicators.'),
        ('Can it run on the target laptop?', 'Concurrent Gazebo and perception benchmarks in FP32, FP16, and INT8 where valid.', 'Frame rate, p95 latency, CPU and GPU use, peak memory, and missed control periods.'),
    ], widths=[1.55, 3.25, 2.20], font_size=8.3)
    keep_table_rows(t)
    add_heading(doc, 'Planned Progression')
    for text in [
        'Stage 1  Reproduce the camera measurement model in Gazebo and compare it with simulator ground truth.',
        'Stage 2  Integrate the estimator, controller, and recovery states for one leader and one follower.',
        'Stage 3  Add followers one at a time and measure how spacing and steering errors propagate through the chain.',
        'Stage 4  Add warehouse aisle geometry, turns, occlusions, and controlled disturbances.',
        'Stage 5  Optimize inference for the RTX 3050 Ti and freeze a reproducible evidence package.',
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run('• ' + text), size=10.3)

    risks_heading = add_heading(doc, '13 Main Risks and Mitigations')
    risks_heading.paragraph_format.page_break_before = True
    t = add_table(doc, ['Risk', 'Effect on the application', 'Planned mitigation'], [
        ('Corner occlusion', 'A follower can lose its predecessor during a sharp turn.', 'Limit leader curvature and speed, model visibility, use bounded prediction, and stop after timeout.'),
        ('Error propagation', 'Rear followers may amplify spacing or steering disturbances.', 'Measure each follower separately, tune conservative gaps, and cap convoy length by evidence.'),
        ('Follower cuts a corner', 'A cart may approach shelving even if the leader path is clear.', 'Test swept paths for every follower and add local obstacle safety independent of formation control.'),
        ('Leader or network fault', 'The mission path may be lost or convoy status may become inconsistent.', 'Propagate faults, stop followers independently, and require supervised recovery.'),
        ('Perception ambiguity', 'Lighting, similar objects, blur, or partial visibility may corrupt the estimate.', 'Use confidence and covariance gates, target-specific appearance, adverse testing, and safe stop.'),
        ('Long convoy at intersections', 'The leader may clear a crossing while followers remain exposed.', 'Reserve the entire convoy footprint at the fleet layer or restrict routes during early deployment.'),
    ], widths=[1.40, 2.75, 2.85], font_size=8.2)
    keep_table_rows(t)

    add_heading(doc, '14 Requested Faculty Decisions')
    add_para(doc,
        'I would like Professor Imtiaz to review the following points before I expand the implementation:')
    for text in [
        'Is the flexible material-delivery convoy a sufficiently concrete application for the EE 616 project?',
        'Is the separation between the visual formation layer and an independent industrial safety layer technically appropriate?',
        'Should the minimum final demonstration require one leader and three followers, or should the validated pair remain the core result with multi-follower behavior as an extension?',
        'Is the planned progression from camera calibration to pair control and then chain stability appropriate for the remaining project schedule?',
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.30)
        p.paragraph_format.first_line_indent = Inches(-0.20)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run('• ' + text), size=10.4)
    add_para(doc,
        'My recommended next step is to approve the application framing and begin the measurement-only ROS 2 and Gazebo baseline. I will keep the existing toy simulator as a deterministic reference, then replace its idealized measurements with rendered camera observations and record the difference. I will not expand to learned detection, multi-follower control, or hardware testing until the preceding evidence and documentation have been reviewed.')

    references_heading = add_heading(doc, 'References')
    references_heading.paragraph_format.page_break_before = True
    add_reference(doc, 1, 'DHL. Annual Report 2016. EffiBOT warehouse test and autonomous trolley description.', 'https://group.dhl.com/content/dam/deutschepostdhl/en/media-center/investors/documents/annual-reports/DPDHL_2016_Annual_Report.pdf')
    add_reference(doc, 2, 'Toyota Material Handling. Tow Tractors and Tuggers FAQ. Multi-cart transport, lean movement, and just-in-time delivery.', 'https://www.toyotaforklift.com/resource-library/blog/toyota-products/tow-tractors-tuggers-faq')
    add_reference(doc, 3, 'Zebra Technologies. Autonomous Mobile Robots. Warehouse workflows including picking, post-pick transport, receiving, putaway, and replenishment.', 'https://www.zebra.com/us/en/products/autonomous-mobile-robots.html')
    add_reference(doc, 4, 'DHL. Logistics Trend Radar. Follow-me and point-to-point mobile robot use in logistics.', 'https://www.dhl.com/content/dam/dhl/global/csi/documents/pdf/csi-logistics-trend-radar-6-dhl.pdf')
    add_reference(doc, 5, 'Zebra Technologies. Team Intelligence FAQ. Description of the single-picker follow-me workflow.', 'https://www.zebra.com/gb/en/resource-library/faq/what-is-team-intelligence.html')
    add_reference(doc, 6, 'International Organization for Standardization. ISO 3691-4:2023 Industrial trucks Safety requirements and verification Part 4 Driverless industrial trucks and their systems.', 'https://www.iso.org/standard/83545.html')
    add_reference(doc, 7, 'United States Occupational Safety and Health Administration. Warehousing hazards and solutions.', 'https://www.osha.gov/warehousing/hazards-solutions')
    add_reference(doc, 8, 'Open Robotics. Gazebo Harmonic feature comparison and ROS 2 Jazzy support.', 'https://gazebosim.org/docs/harmonic/comparison/')
    add_reference(doc, 9, 'Open Robotics. ROS 2 integration with Gazebo Harmonic.', 'https://gazebosim.org/docs/harmonic/ros2_integration/')

    for sec in doc.sections:
        for footer in [sec.footer, sec.first_page_footer, sec.even_page_footer]:
            for p in footer.paragraphs:
                if p.text.strip():
                    for run in p.runs:
                        run.text = ''
                    set_run_font(p.add_run('EE 616 Industry Application Rationale | Neeraj Kumar Kanchani'), size=8.3, color=MID)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    props = doc.core_properties
    props.title = 'Industry Application and System Rationale for Vision Only Leader Follower Warehouse Coordination'
    props.subject = 'Flexible material delivery convoy proposal for faculty review'
    props.author = 'Neeraj Kumar Kanchani'
    props.keywords = 'EE 616, warehouse robotics, leader follower, autonomous mobile robot, visual servoing, ROS 2, Gazebo'
    props.comments = 'Draft prepared by Neeraj Kumar Kanchani for review by Professor Masudul Imtiaz'
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
