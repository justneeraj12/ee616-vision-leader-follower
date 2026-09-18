from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
PACKAGE_WORK = ROOT / 'work' / 'professor_package'
SIM_WORK = ROOT / 'work' / 'toy_simulation'
RESULTS = SIM_WORK / 'results'
OUTPUT = ROOT / 'deliverables' / 'Neeraj_Kanchani_EE616_Initial_Toy_Simulation_Results.docx'
REFERENCE = ROOT / 'deliverables' / 'Neeraj_Kanchani_EE616_Project_Update_Package.docx'

sys.path.insert(0, str(PACKAGE_WORK))
from build_package import (  # noqa: E402
    BLACK,
    MID,
    NAVY,
    add_heading,
    add_para,
    add_table,
    clear_document_body,
    configure_styles,
    set_cell_margins,
    set_repeat_table_header,
    set_run_font,
)


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set('name', title)
    doc_pr.set('descr', description)


def add_figure(doc, path: Path, caption: str, alt: str, width: float = 7.0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt_text(shape, path.stem, alt)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(8)
    run = caption_p.add_run(caption)
    set_run_font(run, size=8.5, italic=True, color=MID)


def build() -> None:
    with (RESULTS / 'summary.json').open(encoding='utf-8') as stream:
        summary = json.load(stream)
    med = summary['overall_median']

    doc = Document(REFERENCE)
    clear_document_body(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)

    top = doc.add_paragraph()
    top.paragraph_format.space_before = Pt(88)
    top.paragraph_format.space_after = Pt(4)
    run = top.add_run('EE 616 Preliminary Evidence')
    set_run_font(run, size=21, color=MID)

    title = doc.add_paragraph(style='Title')
    title.add_run('Initial Toy Simulation Results for Vision Only Leader Follower Coordination')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        set_run_font(run, size=29, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(70)
    run = subtitle.add_run('Kinematic baseline for faculty review before ROS 2 and Gazebo implementation')
    set_run_font(run, size=13.5, color='34495E')

    meta = doc.add_table(rows=1, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_repeat_table_header(meta.rows[0])
    labels = [
        ('STATUS', 'Toy baseline completed'),
        ('OWNER', 'Neeraj Kumar Kanchani'),
        ('LAST UPDATED', 'September 16, 2026'),
    ]
    for index, (label, value) in enumerate(labels):
        cell = meta.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.width = Inches(2.35)
        set_cell_margins(cell, 80, 40, 80, 40)
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        set_run_font(p1.add_run(label), size=8.5, bold=True, color=MID)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=10.2, color='24384D')

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_table(
        doc,
        ['Field', 'Details'],
        [
            ('Faculty reviewer', 'Professor Masudul Imtiaz'),
            ('Course', 'EE 616-57, section 9189, four credits'),
            ('Experiment type', 'Deterministic two-dimensional kinematic simulation'),
            ('Trial set', '32 trials; four scenarios; eight recorded random seeds'),
            ('Decision requested', 'Approve the measurement-only Gazebo baseline as the next implementation gate'),
        ],
        widths=[1.55, 5.45],
        font_size=9.2,
    )

    doc.add_page_break()
    add_heading(doc, '1 Result and Interpretation')
    add_para(
        doc,
        'The toy baseline completed all 32 deterministic trials without a collision sample. Median camera-geometry error was '
        f'{med["range_rmse_m"]:.3f} m for range and {med["bearing_mae_deg"]:.3f} degrees for bearing. Median formation error was '
        f'{med["formation_rmse_m"]:.3f} m, and the forced occlusion case reacquired a measurement after '
        f'{summary["occlusion_median_reacquisition_s"]:.3f} s. These values show that the experiment harness, estimator, controller, and logging contract operate together under simplified assumptions. They are not evidence of ROS 2, Gazebo, learned detection, or physical robot performance.'
    )
    add_figure(
        doc,
        RESULTS / 'figure_4_summary.png',
        'Figure 1  Aggregate toy baseline results with proposed project thresholds shown only as references',
        'Four summary cards report median range error, bearing error, formation error, and occlusion reacquisition time for 32 toy trials.',
        width=7.0,
    )
    add_heading(doc, '2 Experiment Design')
    add_table(
        doc,
        ['Element', 'Toy baseline implementation'],
        [
            ('Scenarios', 'Straight motion, S-curve motion, constant-radius turn, and a forced visual occlusion'),
            ('Perception', 'Synthetic bounding-box height and center with pixel noise at a nominal 15 Hz camera rate'),
            ('Estimation', 'World-frame constant-velocity Kalman filter using follower pose and camera-derived range and bearing'),
            ('Control', 'Bounded unicycle range-and-bearing controller at 50 Hz with TRACK, PREDICT, and SAFE_STOP states'),
            ('Evidence', 'Per-trial metrics, representative time series, fixed configuration, recorded seeds, and four figures'),
        ],
        widths=[1.35, 5.65],
        font_size=9.0,
    )

    representative_heading = add_heading(doc, '3 Representative Behavior')
    representative_heading.paragraph_format.page_break_before = True
    add_figure(
        doc,
        RESULTS / 'figure_1_range_baseline.png',
        'Figure 2  True measured and filtered leader range during the representative forced-occlusion trial',
        'Range over 20 seconds. The true and filtered lines remain near 1.5 meters, measured samples contain small noise, and the forced occlusion is marked from 8.0 to 9.17 seconds.',
        width=6.75,
    )
    add_figure(
        doc,
        RESULTS / 'figure_2_formation_tracking.png',
        'Figure 3  Leader and follower world paths during the representative S-curve trial',
        'Two smooth paths show the follower tracking behind the leader on an S-curve while preserving the desired separation.',
        width=6.75,
    )

    occlusion_heading = add_heading(doc, '4 Occlusion Handling and Limitations')
    occlusion_heading.paragraph_format.page_break_before = True
    add_figure(
        doc,
        RESULTS / 'figure_3_occlusion_recovery.png',
        'Figure 4  Controller state through the forced visual occlusion',
        'A time window around the occlusion shows TRACK samples in green and PREDICT samples in orange before tracking resumes.',
        width=6.75,
    )
    add_heading(doc, 'Limits of This Evidence')
    add_para(
        doc,
        'The follower pose is available without an odometry error model, the leader has a known marker height, and the camera model generates bounding-box geometry directly. The simulation does not include detector misses outside the declared occlusion, image artifacts, variable inference latency, wheel slip, actuator dynamics, message delay, or GPU contention. The low errors are therefore expected and should not be compared with real deployment performance.'
    )
    add_heading(doc, 'Recommended Next Gate')
    add_para(
        doc,
        'I recommend reproducing only the calibrated range-and-bearing experiment in Gazebo next. That gate should compare camera-derived measurements against Gazebo ground truth across range, heading, lighting, and off-axis placement. I will not treat the detector, filter, or controller as validated until that evidence is reviewed.'
    )
    add_heading(doc, 'Approval Requested')
    add_para(
        doc,
        'Please confirm whether this toy baseline and experiment structure are sufficient to begin the measurement-only Gazebo implementation. I will preserve the same seeds, configuration manifest, CSV schema, figures, and acceptance criteria so the next result can be compared directly with this baseline.'
    )

    for sec in doc.sections:
        for footer in [sec.footer, sec.first_page_footer, sec.even_page_footer]:
            for paragraph in footer.paragraphs:
                if paragraph.text.strip():
                    for run in paragraph.runs:
                        run.text = ''
                    new_run = paragraph.add_run('EE 616 Toy Simulation Results | Neeraj Kumar Kanchani')
                    set_run_font(new_run, size=8.3, color=MID)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    props = doc.core_properties
    props.title = 'Initial Toy Simulation Results for Vision Only Leader Follower Coordination'
    props.subject = 'Kinematic baseline for faculty review before ROS 2 and Gazebo implementation'
    props.author = 'Neeraj Kumar Kanchani'
    props.keywords = 'EE 616, leader follower, vision, toy simulation, Kalman filter, control'
    props.comments = 'Prepared for review by Professor Masudul Imtiaz'
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
