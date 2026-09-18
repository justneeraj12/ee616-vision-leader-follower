from __future__ import annotations

import csv
import json
import math
import shutil
import sys
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "work" / "technical_explanation_report"
FIGURES = WORK / "figures"
OUTPUT = ROOT / "deliverables" / "Neeraj_Kanchani_EE616_Initial_Technical_Explanation_and_Preliminary_Figures.docx"
REFERENCE = ROOT / "deliverables" / "Neeraj_Kanchani_EE616_Project_Update_Package.docx"
PAIR_SUMMARY = ROOT / "work" / "toy_simulation" / "results" / "summary.json"
WAREHOUSE_SUMMARY = ROOT / "work" / "toy_simulation" / "warehouse_results" / "warehouse_multi_summary.json"
WAREHOUSE_CSV = ROOT / "work" / "toy_simulation" / "warehouse_results" / "warehouse_multi_timeseries.csv"

sys.path.insert(0, str(ROOT / "work" / "professor_package"))
from build_package import (  # noqa: E402
    BLACK,
    MID,
    NAVY,
    add_heading,
    add_hyperlink,
    add_para,
    add_table,
    clear_document_body,
    configure_styles,
    set_cell_margins,
    set_repeat_table_header,
    set_run_font,
)


BLUE = "#2F6F9F"
LIGHT_BLUE = "#EAF2F7"
PALE_BLUE = "#F5F9FC"
TEXT = "#34495E"
ORANGE = "#C77C18"
GREEN = "#2E7D5B"
RED = "#A94442"
NAVY = f"#{NAVY}"


def font(size: int, bold: bool = False):
    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(f"/usr/share/fonts/truetype/dejavu/{name}", size)


def wrap(draw: ImageDraw.ImageDraw, text: str, box_width: int, face) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        proposed = f"{current} {word}".strip()
        if draw.textbbox((0, 0), proposed, font=face)[2] <= box_width or not current:
            current = proposed
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE, width: int = 8) -> None:
    draw.line((*start, *end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 28
    spread = 0.55
    draw.polygon(
        [
            end,
            (end[0] - size * math.cos(angle - spread), end[1] - size * math.sin(angle - spread)),
            (end[0] - size * math.cos(angle + spread), end[1] - size * math.sin(angle + spread)),
        ],
        fill=color,
    )


def labeled_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], title: str, body: str, fill: str = "white", edge: str = NAVY) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=edge, width=5)
    draw.text((x1 + 26, y1 + 22), title, font=font(31, True), fill=NAVY)
    lines = wrap(draw, body, x2 - x1 - 52, font(23))
    draw.multiline_text((x1 + 26, y1 + 74), "\n".join(lines), font=font(23), fill=TEXT, spacing=8)


def make_topology_figure() -> Path:
    path = FIGURES / "figure_1_system_topology.png"
    image = Image.new("RGB", (3000, 1320), PALE_BLUE)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((35, 28, 2965, 145), radius=18, fill=NAVY)
    draw.text((90, 60), "Leader Follower Topology and Evidence Boundary", font=font(48, True), fill="white")

    labeled_box(draw, (80, 250, 690, 540), "Leader robot", "Owns the waypoint route in the research simulation and exposes a stable visual target.", LIGHT_BLUE)
    labeled_box(draw, (900, 250, 1450, 540), "Follower 1", "Uses its forward RGB camera and local odometry to track the leader.")
    labeled_box(draw, (1660, 250, 2210, 540), "Follower 2", "Tracks Follower 1 through the same local perception and control interface.")
    labeled_box(draw, (2420, 250, 2920, 540), "Follower 3", "Tracks Follower 2. Rearward error growth is measured rather than assumed away.")
    arrow(draw, (690, 395), (900, 395))
    arrow(draw, (1450, 395), (1660, 395))
    arrow(draw, (2210, 395), (2420, 395))

    draw.text((90, 650), "Follower data path", font=font(33, True), fill=NAVY)
    stages = [
        ("RGB image", "camera frames"),
        ("Detection", "target box and confidence"),
        ("Measurement", "range and bearing"),
        ("Estimator", "relative state and uncertainty"),
        ("Supervisor", "TRACK PREDICT SAFE STOP"),
        ("Controller", "bounded velocity commands"),
    ]
    left = 80
    box_width = 400
    gap = 82
    for index, (title, body) in enumerate(stages):
        x1 = left + index * (box_width + gap)
        labeled_box(draw, (x1, 715, x1 + box_width, 995), title, body, "white")
        if index < len(stages) - 1:
            arrow(draw, (x1 + box_width, 855), (x1 + box_width + gap, 855), width=7)

    draw.rounded_rectangle((80, 1100, 2920, 1260), radius=20, fill="#FFF6E8", outline=ORANGE, width=4)
    boundary = (
        "Gazebo ground truth is reserved for evaluation. It must not enter detection, estimation, supervision, or control. "
        "The formation layer is not a safety-rated obstacle or person-protection system."
    )
    draw.multiline_text((120, 1135), "\n".join(wrap(draw, boundary, 2760, font(27))), font=font(27), fill=TEXT, spacing=8)
    image.save(path, quality=95)
    return path


def make_geometry_figure() -> Path:
    path = FIGURES / "figure_2_camera_geometry.png"
    image = Image.new("RGB", (3000, 1450), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((35, 28, 2965, 145), radius=18, fill=NAVY)
    draw.text((90, 60), "Camera Geometry Used by the Toy Measurement Model", font=font(48, True), fill="white")

    camera_x, camera_y = 420, 800
    target_x = 1860
    draw.polygon([(270, 690), (510, 690), (590, 800), (510, 910), (270, 910)], fill=LIGHT_BLUE, outline=NAVY)
    draw.ellipse((camera_x - 35, camera_y - 35, camera_x + 35, camera_y + 35), fill=NAVY)
    draw.text((235, 945), "Follower camera", font=font(30, True), fill=NAVY)

    draw.line((camera_x, camera_y, target_x, 430), fill=BLUE, width=7)
    draw.line((camera_x, camera_y, target_x, 1170), fill=BLUE, width=7)
    draw.line((camera_x, camera_y, target_x, camera_y), fill="#8899A6", width=4)
    draw.arc((480, 680, 850, 920), start=315, end=360, fill=ORANGE, width=8)
    draw.text((700, 690), "bearing θ", font=font(31, True), fill=ORANGE)
    draw.line((target_x, 430, target_x, 1170), fill=GREEN, width=18)
    draw.text((target_x + 45, 760), "known target\nheight H", font=font(29, True), fill=GREEN, spacing=8)
    draw.text((1020, 825), "relative range r", font=font(31, True), fill=BLUE)

    draw.rounded_rectangle((2120, 250, 2890, 1215), radius=24, fill=PALE_BLUE, outline=NAVY, width=5)
    draw.text((2170, 300), "Image measurement", font=font(34, True), fill=NAVY)
    draw.rectangle((2250, 430, 2760, 980), outline="#60758A", width=6)
    draw.line((2505, 430, 2505, 980), fill="#AAB5BF", width=3)
    draw.rectangle((2370, 555, 2665, 900), outline=GREEN, width=10)
    draw.line((2517, 555, 2517, 900), fill=GREEN, width=4)
    draw.text((2580, 915), "box height h", font=font(27, True), fill=GREEN)
    draw.text((2230, 1015), "box center u", font=font(27, True), fill=ORANGE)
    draw.text((2190, 1100), "principal point cₓ", font=font(25), fill=TEXT)

    formula_y = 1245
    draw.rounded_rectangle((90, formula_y, 2910, 1395), radius=20, fill=LIGHT_BLUE, outline=BLUE, width=4)
    draw.text((170, formula_y + 38), "Range estimate  r = f × H ÷ h", font=font(35, True), fill=NAVY)
    draw.text((1510, formula_y + 38), "Bearing estimate  θ = atan((u − cₓ) ÷ f)", font=font(35, True), fill=NAVY)
    image.save(path, quality=95)
    return path


def make_state_figure() -> Path:
    path = FIGURES / "figure_3_state_supervisor.png"
    image = Image.new("RGB", (3000, 1250), PALE_BLUE)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((35, 28, 2965, 145), radius=18, fill=NAVY)
    draw.text((90, 60), "Follower Estimation Control and Recovery Logic", font=font(48, True), fill="white")

    labeled_box(draw, (100, 245, 670, 510), "Measurement", "Camera-derived range and bearing arrive with timestamps and validity flags.", LIGHT_BLUE)
    labeled_box(draw, (890, 245, 1500, 510), "Kalman filter", "Predict the predecessor state between image updates and correct it when a valid measurement arrives.")
    labeled_box(draw, (1720, 245, 2320, 510), "Bounded control", "Use estimated gap and bearing to command clipped linear and angular speed.")
    labeled_box(draw, (2540, 245, 2900, 510), "Robot", "Apply velocity commands in the toy kinematic model.", LIGHT_BLUE)
    arrow(draw, (670, 378), (890, 378))
    arrow(draw, (1500, 378), (1720, 378))
    arrow(draw, (2320, 378), (2540, 378))

    draw.text((100, 630), "Supervisory states", font=font(35, True), fill=NAVY)
    states = [
        ((120, 720, 850, 1040), "TRACK", "A recent valid measurement updates the estimate. Normal bounded control is allowed.", GREEN),
        ((1135, 720, 1865, 1040), "PREDICT", "Detection is temporarily absent. The filter predicts for no more than the configured timeout.", ORANGE),
        ((2150, 720, 2880, 1040), "SAFE STOP", "The estimate is unavailable or too old. Linear and angular commands are set to zero.", RED),
    ]
    for xy, title, body, color in states:
        x1, y1, x2, y2 = xy
        draw.rounded_rectangle(xy, radius=28, fill="white", outline=color, width=8)
        draw.text((x1 + 32, y1 + 28), title, font=font(38, True), fill=color)
        draw.multiline_text((x1 + 32, y1 + 100), "\n".join(wrap(draw, body, x2 - x1 - 64, font(25))), font=font(25), fill=TEXT, spacing=9)
    arrow(draw, (850, 880), (1135, 880), ORANGE)
    arrow(draw, (1865, 880), (2150, 880), RED)
    arrow(draw, (1135, 970), (850, 970), GREEN)
    draw.text((890, 1085), "valid detection returns", font=font(24), fill=GREEN)
    draw.text((125, 1120), "The timeout is a research supervisor rule, not a certified stopping-distance guarantee.", font=font(27, True), fill=TEXT)
    image.save(path, quality=95)
    return path


def make_results_figure(pair: dict, warehouse: dict) -> Path:
    path = FIGURES / "figure_4_preliminary_results.png"
    with WAREHOUSE_CSV.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    stride = 8
    sample = rows[::stride]
    colors = ["#0B3153", "#2F6F9F", "#2E7D5B", "#C77C18"]
    image = Image.new("RGB", (3000, 1600), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((35, 28, 2965, 145), radius=18, fill=NAVY)
    draw.text((90, 60), "Preliminary Toy Results and Rearward Error Growth", font=font(48, True), fill="white")

    path_box = (90, 235, 1480, 1510)
    bar_box = (1600, 235, 2910, 770)
    time_box = (1600, 900, 2910, 1510)
    for box in (path_box, bar_box, time_box):
        draw.rounded_rectangle(box, radius=18, fill=PALE_BLUE, outline="#B7C5D1", width=4)

    draw.text((135, 275), "Recorded 86 second warehouse path", font=font(31, True), fill=NAVY)
    left, top, right, bottom = 155, 350, 1415, 1390
    def map_path(x: float, y: float) -> tuple[int, int]:
        return int(left + (x / 16.5) * (right - left)), int(bottom - (y / 10.2) * (bottom - top))
    shelves = (
        (1.5, 2.2, 6.8, 3.4), (9.2, 2.2, 14.5, 3.4),
        (1.5, 4.9, 6.8, 6.1), (9.2, 4.9, 14.5, 6.1),
        (1.5, 7.6, 6.8, 8.6), (9.2, 7.6, 14.5, 8.6),
    )
    for x1, y1, x2, y2 in shelves:
        px1, py2 = map_path(x1, y1)
        px2, py1 = map_path(x2, y2)
        draw.rectangle((px1, py1, px2, py2), fill="#D9E1E8", outline="#60758A", width=3)
    labels = ["Leader", "Follower 1", "Follower 2", "Follower 3"]
    for index, label in enumerate(labels):
        points = [map_path(float(row[f"r{index}_x_m"]), float(row[f"r{index}_y_m"])) for row in sample]
        draw.line(points, fill=colors[index], width=7)
        key_x = 180 + index * 300
        draw.line((key_x, 1445, key_x + 55, 1445), fill=colors[index], width=8)
        draw.text((key_x + 68, 1427), label, font=font(21), fill=TEXT)

    rmse = [warehouse["spacing_rmse_m"][f"follower_{index}"] for index in range(1, 4)]
    draw.text((1645, 275), "Spacing RMSE grows toward the rear", font=font(31, True), fill=NAVY)
    bar_left, bar_top, bar_right, bar_bottom = 1690, 365, 2835, 675
    draw.line((bar_left, bar_bottom, bar_right, bar_bottom), fill="#60758A", width=3)
    draw.line((bar_left, bar_top, bar_left, bar_bottom), fill="#60758A", width=3)
    for index, value in enumerate(rmse):
        center = bar_left + 220 + index * 355
        height = int((value / 0.24) * (bar_bottom - bar_top))
        draw.rounded_rectangle((center - 90, bar_bottom - height, center + 90, bar_bottom), radius=14, fill=colors[index + 1])
        draw.text((center - 52, bar_bottom - height - 45), f"{value:.3f}", font=font(24, True), fill=NAVY)
        draw.text((center - 67, bar_bottom + 18), f"Follower {index + 1}", font=font(21), fill=TEXT)

    draw.text((1645, 940), "Spacing error relative to the 0.8 m command", font=font(31, True), fill=NAVY)
    graph_left, graph_top, graph_right, graph_bottom = 1690, 1030, 2835, 1405
    draw.line((graph_left, graph_bottom, graph_right, graph_bottom), fill="#60758A", width=3)
    draw.line((graph_left, graph_top, graph_left, graph_bottom), fill="#60758A", width=3)
    zero_y = int(graph_top + (0.8 / 1.6) * (graph_bottom - graph_top))
    draw.line((graph_left, zero_y, graph_right, zero_y), fill="#9AA8B3", width=2)
    def map_error(time_s: float, error: float) -> tuple[int, int]:
        x = graph_left + (time_s / 86.0) * (graph_right - graph_left)
        clipped = max(-0.8, min(0.8, error))
        y = graph_top + ((0.8 - clipped) / 1.6) * (graph_bottom - graph_top)
        return int(x), int(y)
    for index in range(1, 4):
        points = [map_error(float(row["time_s"]), float(row[f"r{index}_true_range_m"]) - 0.8) for row in sample]
        draw.line(points, fill=colors[index], width=4)
        key_x = 1780 + (index - 1) * 260
        draw.line((key_x, 1460, key_x + 45, 1460), fill=colors[index], width=6)
        draw.text((key_x + 55, 1443), f"F{index}", font=font(20), fill=TEXT)
    draw.text((2790, 1420), "86 s", font=font(20), fill=TEXT)
    draw.text((1608, 1015), "+0.8", font=font(18), fill=TEXT)
    draw.text((1625, zero_y - 12), "0", font=font(18), fill=TEXT)
    draw.text((1608, 1385), "−0.8", font=font(18), fill=TEXT)

    med = pair["overall_median"]
    footer = "Toy kinematic evidence only  |  32 pair trials  |  range RMSE %.3f m  |  bearing MAE %.3f°  |  warehouse collision samples %d" % (
        med["range_rmse_m"], med["bearing_mae_deg"], warehouse["collision_samples"]
    )
    draw.rounded_rectangle((90, 1530, 2910, 1585), radius=14, fill=LIGHT_BLUE)
    draw.text((140, 1542), footer, font=font(22, True), fill=NAVY)
    image.save(path, quality=95)
    return path


def keep_table_rows(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_figure(doc, path: Path, caption: str, alt: str, width: float = 7.0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("name", path.stem)
    shape._inline.docPr.set("descr", alt)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(8)
    set_run_font(caption_p.add_run(caption), size=8.5, italic=True, color=MID)


def add_reference(doc, number: int, citation: str, url: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(5)
    set_run_font(paragraph.add_run(f"[{number}] {citation} "), size=8.8)
    add_hyperlink(paragraph, "Source link", url)


def build() -> Path:
    WORK.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    with PAIR_SUMMARY.open(encoding="utf-8") as stream:
        pair = json.load(stream)
    with WAREHOUSE_SUMMARY.open(encoding="utf-8") as stream:
        warehouse = json.load(stream)
    med = pair["overall_median"]

    figures = [
        make_topology_figure(),
        make_geometry_figure(),
        make_state_figure(),
        make_results_figure(pair, warehouse),
    ]

    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_document_body(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)

    top = doc.add_paragraph()
    top.paragraph_format.space_before = Pt(78)
    top.paragraph_format.space_after = Pt(4)
    set_run_font(top.add_run("EE 616 Technical Review"), size=21, color=MID)

    title = doc.add_paragraph(style="Title")
    title.add_run("Initial Technical Explanation and Preliminary Figures for Vision Only Leader Follower Coordination")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(12)
    for run in title.runs:
        set_run_font(run, size=28, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(54)
    set_run_font(subtitle.add_run("Technical concept and toy kinematic evidence for faculty review"), size=13.5, color=TEXT.lstrip("#"))

    meta = doc.add_table(rows=1, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_repeat_table_header(meta.rows[0])
    for index, (label, value) in enumerate([
        ("STATUS", "Draft for faculty review"),
        ("AUTHOR", "Neeraj Kumar Kanchani"),
        ("UPDATED", "September 18, 2026"),
    ]):
        cell = meta.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.width = Inches(2.35)
        set_cell_margins(cell, 85, 45, 85, 45)
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        set_run_font(p1.add_run(label), size=8.5, bold=True, color=MID)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=10.1, color="24384D")

    doc.add_paragraph().paragraph_format.space_after = Pt(5)
    table = add_table(doc, ["Field", "Details"], [
        ("Faculty reviewer", "Professor Masudul Imtiaz"),
        ("Course", "EE 616-57, section 9189, four credits"),
        ("Application", "Flexible autonomous material-delivery convoy for warehouse and manufacturing aisles"),
        ("Current evidence", "Deterministic two-dimensional toy simulation"),
        ("Decision requested", "Confirm the technical direction and approve environment validation as the next implementation gate"),
    ], widths=[1.55, 5.45], font_size=9.0)
    keep_table_rows(table)

    doc.add_page_break()
    add_heading(doc, "1 Technical Purpose and Review Request")
    add_para(doc,
        "I propose a vision-only leader-follower coordination system for a flexible warehouse material-delivery convoy. The leader owns the route. Each independently driven follower observes the robot immediately ahead with a forward RGB camera, combines that observation with its own wheel odometry, and commands bounded motion to maintain a specified gap. My technical question is whether this local perception and control structure can preserve useful spacing through turns and short losses of visual contact on constrained computing hardware.")
    add_para(doc,
        "The current evidence is a deterministic two-dimensional Python model. It demonstrates the planned predecessor-following topology, camera-geometry calculation, state estimation, bounded control, recovery-state logic, synchronized logging, and repeatable figures. It does not demonstrate ROS 2 timing, Gazebo physics, learned image detection, wheel slip, network behavior, physical robots, or industrial safety.")
    add_para(doc,
        "I am asking for feedback on the technical decomposition, the proposed measurement experiment, and the decision to validate one leader and one follower before extending the chain. If this direction is suitable, my next implementation request will be limited to environment and reproducibility validation for ROS 2 Jazzy and Gazebo Harmonic.")

    add_heading(doc, "2 Verified State")
    table = add_table(doc, ["Item", "Verified condition", "Meaning"], [
        ("Automated tests", "Seven tests pass", "The recorded toy functions remain deterministic and finite."),
        ("Pair experiment", "32 trials across four scenarios and eight seeds", "The experiment contract is repeatable under simplified assumptions."),
        ("Warehouse run", "One leader, three followers, 86 seconds, seven route segments", "The chain topology and logging operate in the toy maze."),
        ("Collision record", "Zero toy-model collision samples", "No collision was recorded by the center and radius checks; this is not a safety claim."),
        ("ROS and Gazebo", "Complete working stack not verified", "No ROS 2 or Gazebo performance result is claimed."),
    ], widths=[1.25, 2.45, 3.30], font_size=8.6)
    keep_table_rows(table)

    doc.add_page_break()
    add_heading(doc, "3 System Topology and Experimental Boundary")
    add_figure(doc, figures[0], "Figure 1  Proposed predecessor-following topology and the boundary between control inputs and evaluation data", "A leader is followed by three followers. Each follower uses an RGB image, target detection, relative measurement, an estimator, a supervisor, and bounded control. Gazebo ground truth is shown outside the control path.", 7.0)
    add_para(doc,
        "Follower 1 tracks the leader, Follower 2 tracks Follower 1, and later robots repeat the same interface. This local relationship avoids requiring a rear robot to see the leader through the entire convoy. It also creates a measurable risk: spacing and turning error can grow toward the rear. I will therefore add followers one at a time only after the one-pair experiment is validated.")
    add_para(doc,
        "ROS 2 topics are suitable for continuous sensor and state streams, and namespaces provide the naming boundary needed to instantiate the same follower pipeline more than once. In the proposed implementation, each robot will have an isolated namespace, and message timestamps will be logged with the configuration and trial identifier. [1]")

    doc.add_page_break()
    add_heading(doc, "4 Camera Only Range and Bearing Measurement")
    add_figure(doc, figures[1], "Figure 2  Monocular range and bearing calculation used by the toy camera model", "A pinhole camera observes a target with known physical height. Range is estimated from focal length, target height, and bounding-box height. Bearing is estimated from horizontal pixel offset.", 7.0)
    add_para(doc,
        "The toy model assumes a known target height H, focal length f in pixels, measured bounding-box height h, horizontal box center u, and camera principal point cₓ. It estimates range as r = f × H ÷ h and bearing as θ = atan((u − cₓ) ÷ f). Pixel noise is added before those calculations. A later image detector would supply the box and confidence, but the geometric calculation would remain independently testable.")
    add_para(doc,
        "This method is intentionally narrow. Range error will increase when the detected height is small, clipped, tilted, or inconsistent with the assumed target geometry. The Gazebo measurement baseline must therefore sweep range, off-axis placement, heading, lighting, and partial visibility. Gazebo Harmonic provides simulated camera sensors and configurable sensor update behavior, but the project must measure the resulting error rather than infer accuracy from simulator availability. [2]")

    doc.add_page_break()
    add_heading(doc, "5 Estimation Control and Recovery")
    add_figure(doc, figures[2], "Figure 3  Estimation, bounded control, and the TRACK, PREDICT, and SAFE STOP supervisor", "A camera measurement updates a Kalman filter. The controller uses the estimate while the supervisor chooses TRACK, PREDICT, or SAFE STOP based on measurement age.", 7.0)
    add_para(doc,
        "The current filter estimates the predecessor's two-dimensional position and velocity. It predicts at every 0.02 second step and corrects the estimate when a camera sample is available. The toy implementation transforms a relative camera measurement into the world frame using an ideal follower pose. That ideal pose is a major simplification; the ROS 2 version must use local wheel odometry and track covariance and timestamp age explicitly.")
    add_para(doc,
        "The controller combines estimated forward target speed with proportional range and bearing terms, then clips linear and angular commands. TRACK uses a recent measurement. PREDICT allows the filter to bridge a short gap. SAFE STOP sets both commands to zero after the configured timeout or before the filter is initialized. These states express deterministic failure behavior, but they do not establish certified braking distance or person safety.")

    doc.add_page_break()
    add_heading(doc, "6 Preliminary Toy Evidence")
    add_figure(doc, figures[3], "Figure 4  Recorded toy trajectories, spacing error, and error growth toward the rear of the convoy", "The left panel shows one leader and three follower trajectories through six shelf blocks. The right panels show spacing RMSE by follower and spacing error over time.", 7.0)
    add_para(doc,
        f"The single-pair baseline contains 32 trials covering straight motion, an S-curve, a constant-radius turn, and a forced visual occlusion. Across those trials, the median synthetic camera range RMSE was {med['range_rmse_m']:.3f} m, bearing MAE was {med['bearing_mae_deg']:.3f} degrees, formation RMSE was {med['formation_rmse_m']:.3f} m, and median reacquisition after the forced occlusion was {pair['occlusion_median_reacquisition_s']:.3f} s. No collision sample was recorded in that toy experiment.")
    add_para(doc,
        f"The warehouse demonstration used one leader and three followers for {warehouse['duration_s']:.0f} simulated seconds. It completed {warehouse['completed_waypoints']} route segments and recorded zero collision samples in the toy collision model. Spacing RMSE increased from {warehouse['spacing_rmse_m']['follower_1']:.3f} m for Follower 1 to {warehouse['spacing_rmse_m']['follower_2']:.3f} m for Follower 2 and {warehouse['spacing_rmse_m']['follower_3']:.3f} m for Follower 3. I interpret this as an early indication that rearward error propagation and corner behavior require explicit measurement in the later multi-follower gate.")

    doc.add_page_break()
    add_heading(doc, "7 What the Evidence Does Not Prove")
    table = add_table(doc, ["Missing condition", "Why it matters", "Required later evidence"], [
        ("Rendered camera frames and learned detection", "The toy model generates boxes directly and omits image artifacts and false detections.", "Gazebo images, labeled detections, confidence, misses, and error against evaluation ground truth."),
        ("Robot physics and odometry error", "The model omits wheel slip, actuator lag, contact dynamics, and odometry drift.", "Gazebo motion logs and controlled disturbances after the measurement gate."),
        ("ROS 2 timing and namespaces", "The current code does not exercise middleware timing, queueing, QoS, or namespace isolation.", "Timestamp audits, topic graphs, launch tests, and dropped-frame measurements."),
        ("Edge inference resource use", "Synthetic geometry does not measure detector latency or GPU memory.", "FP32 baseline first, followed by justified optimization on the target laptop."),
        ("Industrial safety", "A formation controller is not a safety-rated obstacle or person-protection system.", "Independent safety architecture, verified stopping behavior, site risk assessment, commissioning, and applicable standards work."),
    ], widths=[1.55, 2.65, 2.80], font_size=8.4)
    keep_table_rows(table)
    add_para(doc,
        "ISO 3691-4:2023 addresses safety requirements and verification for driverless industrial trucks and their systems. My simulation will not be presented as compliance evidence. I will keep formation coordination separate from independent safety functions throughout the report and implementation. [3]")

    add_heading(doc, "8 Recommended Next Technical Gate")
    add_para(doc,
        "I recommend validating the development environment before building the camera experiment. The present machine contains part of a ROS 2 Jazzy installation, but a working ros2 command and Gazebo gz command have not been verified. The environment gate should record exact package versions, shell setup, a namespaced publisher and subscriber smoke test, a minimal Gazebo camera world, CPU and GPU availability, and reproducible setup instructions. Gazebo Harmonic is the intended long-term release and lists Ubuntu Noble on amd64 as a supported platform. [4]")
    add_para(doc,
        "After that gate passes, I recommend a measurement-only experiment. The follower camera should observe a known predecessor target across a controlled pose grid. Estimated range and bearing will be compared with synchronized Gazebo ground truth used only by the evaluator. The detector, filter, and controller should remain outside that gate so a measurement error can be diagnosed without closed-loop behavior masking it.")

    questions_heading = add_heading(doc, "9 Questions for Faculty Review")
    questions_heading.paragraph_format.page_break_before = True
    table = add_table(doc, ["Question", "Decision affected"], [
        ("Is the one-pair measurement experiment an appropriate minimum scientific unit before closed-loop control?", "Determines whether implementation proceeds in the staged order proposed here."),
        ("Is a known visual target acceptable for the first Gazebo baseline before introducing a learned detector?", "Determines whether camera geometry can be isolated from detector performance."),
        ("Are the proposed error, timing, and resource measurements sufficient for the first review?", "Determines the acceptance criteria recorded before implementation."),
        ("Should the final demonstration target one leader and three followers if the pair and incremental-chain evidence pass?", "Determines the intended December demonstration scale."),
    ], widths=[4.55, 2.45], font_size=8.6)
    keep_table_rows(table)

    add_heading(doc, "References")
    add_reference(doc, 1, "ROS 2 Jazzy documentation. Interfaces for topics, services, and actions.", "https://docs.ros.org/en/jazzy/How-To-Guides/Topics-Services-Actions.html")
    add_reference(doc, 2, "Gazebo Harmonic documentation. Sensors tutorial and sensor configuration.", "https://gazebosim.org/docs/harmonic/sensors/")
    add_reference(doc, 3, "International Organization for Standardization. ISO 3691-4:2023 Industrial trucks safety requirements for driverless industrial trucks and their systems.", "https://www.iso.org/standard/83545.html")
    add_reference(doc, 4, "Gazebo Harmonic documentation. Installation, release libraries, and supported platforms.", "https://gazebosim.org/docs/harmonic/install/")

    add_heading(doc, "Project Evidence Files")
    add_para(doc,
        "The numerical results in this report come from work/toy_simulation/results/summary.json, work/toy_simulation/results/trial_metrics.csv, work/toy_simulation/warehouse_results/warehouse_multi_summary.json, and work/toy_simulation/warehouse_results/warehouse_multi_timeseries.csv. The source implementation and automated tests are stored under work/toy_simulation/src and work/toy_simulation/tests.")

    for sec in doc.sections:
        for footer in [sec.footer, sec.first_page_footer, sec.even_page_footer]:
            for paragraph in footer.paragraphs:
                if paragraph.text.strip():
                    for run in paragraph.runs:
                        run.text = ""
                    set_run_font(paragraph.add_run("EE 616 Technical Review | Neeraj Kumar Kanchani"), size=8.3, color=MID)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    properties = doc.core_properties
    properties.title = "Initial Technical Explanation and Preliminary Figures for Vision Only Leader Follower Coordination"
    properties.subject = "Technical concept and toy kinematic evidence for faculty review"
    properties.author = "Neeraj Kumar Kanchani"
    properties.keywords = "EE 616, leader follower, vision, camera geometry, Kalman filter, control, preliminary evidence"
    properties.comments = "Draft prepared for review by Professor Masudul Imtiaz"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
